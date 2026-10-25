"""Export one trip worksheet of the original Google Sheets ledger as an anonymized test fixture.

Usage: uv run python scripts/export_sheet_fixture.py <workbook.xlsx> <sheet name> <output.json>

Names become "Person A", "Person B", ... and descriptions and dates are dropped, so the
fixture keeps only the numbers the settlement engine needs plus the balances the sheet computed.
"""

import json
import string
import sys
from decimal import Decimal

import openpyxl

FIRST_ROW = 4
COL = {"payer": 4, "beneficiaries": 5, "category": 7, "amount": 9, "currency": 10}
COL_BENEFICIARY_NAME, COL_PAID_BY, COL_BALANCE = 16, 17, 44


def main(path: str, sheet: str, out: str) -> None:
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    rate = Decimal(str(ws["M2"].value))
    foreign = str(ws["L2"].value).split()[0]

    names: list[str] = []
    paid_by: dict[str, list[str]] = {}
    expected: dict[str, float] = {}
    row = FIRST_ROW
    while ws.cell(row, COL_BENEFICIARY_NAME).value:
        name = ws.cell(row, COL_BENEFICIARY_NAME).value.strip()
        names.append(name)
        sponsors = ws.cell(row, COL_PAID_BY).value
        paid_by[name] = [s.strip() for s in sponsors.split(",")] if sponsors else []
        expected[name] = round(float(ws.cell(row, COL_BALANCE).value or 0), 2)
        row += 1
    alias = {name: f"Person {string.ascii_uppercase[i]}" for i, name in enumerate(names)}

    expenses = []
    for r in range(FIRST_ROW, ws.max_row + 1):
        payer = ws.cell(r, COL["payer"]).value
        if not payer:
            continue
        expenses.append(
            {
                "payer": alias[payer.strip()],
                "beneficiaries": [
                    alias[b.strip()] for b in ws.cell(r, COL["beneficiaries"]).value.split(",")
                ],
                "category": ws.cell(r, COL["category"]).value,
                "amount": str(ws.cell(r, COL["amount"]).value or 0),
                "currency": ws.cell(r, COL["currency"]).value,
            }
        )

    fixture = {
        "base_currency": "USD",
        "rates": {foreign: str(rate)},
        "participants": [alias[n] for n in names],
        "sponsors": {alias[n]: [alias[s] for s in paid_by[n]] for n in names if paid_by[n]},
        "expenses": expenses,
        "expected_balances": {alias[n]: expected[n] for n in names},
    }
    with open(out, "w") as f:
        json.dump(fixture, f, indent=1, ensure_ascii=False)


if __name__ == "__main__":
    main(*sys.argv[1:4])
